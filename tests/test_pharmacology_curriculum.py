"""Offline contract checks for the source-verified curriculum; no bot/state import."""
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from docx import Document
from scripts.course_automation.source_text import docx_text
from scripts.course_automation import build_pharmacology


def main():
    data = json.loads((ROOT / 'course_sources/pharmacology/curriculum.json').read_text(encoding='utf-8'))
    lessons = [lesson for module in data['modules'] for lesson in module['lessons']]
    assert [lesson['number'] for lesson in lessons] == list(range(1, 50))
    controls = {lesson['assessment']: lesson['number'] for lesson in lessons if 'assessment' in lesson}
    assert controls == dict(zip([f'control_{i}' for i in range(1, 7)], [7, 13, 21, 31, 41, 49]))
    assert [m['after_lesson'] for m in data['milestones']] == [24, 49]
    assert len(data['modules']) == 5
    assert len(data['navigation']) == 5
    assert not data['publication_rules']['unreviewed_ocr_in_student_content']
    assert not data['publication_rules']['unreviewed_ocr_in_ai_knowledge']
    assert callable(build_pharmacology.main)
    with tempfile.TemporaryDirectory(prefix='pharma_source_test_') as folder:
        path = Path(folder) / 'ordered.docx'
        doc = Document()
        doc.add_paragraph('Перед таблицей')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Данные таблицы'
        doc.add_paragraph('После таблицы')
        doc.save(path)
        text = docx_text(path)
        assert text.index('Перед') < text.index('Данные') < text.index('После')
    print('PHARMACOLOGY CURRICULUM: 49 classes / 6 controls / 5 modules; extraction order OK')


if __name__ == '__main__':
    main()
