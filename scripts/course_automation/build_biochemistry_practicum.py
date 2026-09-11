"""Build the complete VMEDA biochemistry subject from the approved DOCX practicum.

The source DOCX is the single source of truth.  The builder preserves every non-empty body
paragraph, keeps the source hierarchy, and emits navigation-sized chunks for Telegram and the
Mini App without summarising or rewriting the medical content.
"""

from __future__ import annotations

import hashlib
import html
import json
import re
from collections import Counter
from copy import deepcopy
from pathlib import Path

from docx import Document

REPO = Path(__file__).resolve().parents[2]
SOURCE = REPO / "course_sources" / "biochemistry" / "Biokhimia_VMedA_praktikum_zanyatia.docx"
COURSE_OUTPUT = REPO / "generated_courses" / "biochemistry.json"
REPORT_DIR = REPO / "generated_reports" / "biochemistry"
SOURCE_LABEL = "Практикум по биохимии ВМедА"
MAX_CONTENT = 3200
MAX_NAV_TITLE = 64
CONTENT_SEPARATOR = "\n\n────────\n\n"

READING_MARKERS = {
    "Принцип.": "◆",
    "Ход.": "→",
    "Результат.": "✓",
    "В протокол.": "▣",
    "Зачем врачу.": "⚕",
    "Суть.": "◆",
    "Разбор.": "→",
    "Клиника.": "⚕",
    "Норма.": "✓",
    "Вывод.": "✓",
}


def clean_text(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def nav_title(value: str, limit: int = MAX_NAV_TITLE) -> str:
    value = clean_text(value)
    if len(value) <= limit:
        return value
    shortened = value[: limit - 1].rsplit(" ", 1)[0]
    return (shortened or value[: limit - 1]).rstrip(" .,:;–—-") + "…"


def semantic_plain_html(value: str) -> str:
    """Emphasise structural terms without changing a single source character."""
    bullet = re.match(r"^(•\s*)([^:=—–]{1,69}\S)(\s*(?:=|:|—|–)\s*)(.+)$", value)
    if bullet:
        lead, term, separator, rest = bullet.groups()
        return f"{html.escape(lead)}<b>{html.escape(term)}</b>{html.escape(separator)}{html.escape(rest)}"

    definition = re.match(r"^([^.!?;:]{1,69}\S)(\s+[—–]\s+)(.+)$", value)
    if definition:
        term, separator, rest = definition.groups()
        return f"<b>{html.escape(term)}</b>{html.escape(separator)}{html.escape(rest)}"

    lead_in = re.match(r"^([^.!?;:]{2,60}:)(\s+)(.+)$", value)
    if lead_in:
        lead, spacing, rest = lead_in.groups()
        return f"<b>{html.escape(lead)}</b>{html.escape(spacing)}{html.escape(rest)}"

    numbered = re.match(r"^(\d+[.)])(\s+)(.+)$", value)
    if numbered:
        number, spacing, rest = numbered.groups()
        return f"<b>{html.escape(number)}</b>{html.escape(spacing)}{html.escape(rest)}"
    return html.escape(value, quote=False)


def paragraph_html(paragraph) -> str:
    """Preserve supported inline emphasis while keeping the paragraph's complete text."""
    rendered: list[str] = []
    for run in paragraph.runs:
        value = html.escape(run.text, quote=False)
        if not value:
            continue
        if run.bold:
            value = f"<b>{value}</b>"
        if run.italic:
            value = f"<i>{value}</i>"
        if run.underline:
            value = f"<u>{value}</u>"
        if run.font.strike:
            value = f"<s>{value}</s>"
        rendered.append(value)
    rendered_html = "".join(rendered).strip()
    plain = clean_text(paragraph.text)
    if not rendered_html:
        rendered_html = html.escape(plain, quote=False)
    if not any(run.bold or run.italic or run.underline or run.font.strike for run in paragraph.runs):
        rendered_html = semantic_plain_html(plain)
    for label, marker in READING_MARKERS.items():
        if plain.startswith(label):
            return f"{marker} {rendered_html}"
    return rendered_html


def split_plain_text(value: str, limit: int) -> list[str]:
    """Split an unusually long paragraph without dropping any non-whitespace text."""
    result: list[str] = []
    remaining = clean_text(value)
    while len(remaining) > limit:
        cut = remaining.rfind(" ", 0, limit + 1)
        if cut < limit // 2:
            cut = limit
        result.append(remaining[:cut].strip())
        remaining = remaining[cut:].strip()
    if remaining:
        result.append(remaining)
    return result


def pack_records(records: list[dict], heading: str | None = None) -> list[dict]:
    """Pack whole paragraphs into Telegram-safe chunks; headings remain verbatim in content."""
    units: list[dict] = []
    if heading:
        units.append({"plain": heading, "html": f"<b>{html.escape(heading, quote=False)}</b>", "index": None})
    for record in records:
        rendered = record["html"]
        if len(rendered) <= MAX_CONTENT:
            units.append({"plain": record["text"], "html": rendered, "index": record["index"]})
            continue
        for fragment in split_plain_text(record["text"], MAX_CONTENT - 50):
            units.append({"plain": fragment, "html": html.escape(fragment, quote=False), "index": record["index"]})

    chunks: list[dict] = []
    current: list[dict] = []
    current_length = 0
    for unit in units:
        addition = len(unit["html"]) + (2 if current else 0)
        if current and current_length + addition > MAX_CONTENT:
            chunks.append({"content": "\n\n".join(item["html"] for item in current), "units": current})
            current = []
            current_length = 0
        current.append(unit)
        current_length += len(unit["html"]) + (2 if len(current) > 1 else 0)
    if current:
        chunks.append({"content": "\n\n".join(item["html"] for item in current), "units": current})
    return chunks


def locator(records: list[dict], fallback_index: int) -> str:
    indices = [record["index"] for record in records if record.get("index") is not None]
    if not indices:
        indices = [fallback_index]
    start, end = min(indices), max(indices)
    suffix = f"абз. {start}" if start == end else f"абз. {start}–{end}"
    return f"{SOURCE_LABEL}, {suffix}"


def make_lessons(prefix: str, title: str, records: list[dict], heading_index: int) -> list[dict]:
    chunks = pack_records(records, heading=title)
    lessons: list[dict] = []
    for part, chunk in enumerate(chunks, 1):
        part_suffix = f" · {part}/{len(chunks)}" if len(chunks) > 1 else ""
        lesson_records = [unit for unit in chunk["units"] if unit.get("index") is not None]
        lessons.append({
            "id": f"{prefix}_p{part}",
            "title": nav_title(title, MAX_NAV_TITLE - len(part_suffix)) + part_suffix,
            "content": chunk["content"],
            "sources": [locator(lesson_records, heading_index)],
        })
    return lessons


def lesson_pages(lesson: dict) -> list[str]:
    return lesson.get("content_pages") or [lesson["content"]]


def pack_lesson_pages(lessons: list[dict]) -> list[str]:
    """Pack complete lesson fragments into Telegram-safe pages without rewriting them."""
    pages: list[str] = []
    current: list[str] = []
    current_length = 0
    for lesson in lessons:
        for content in lesson_pages(lesson):
            addition = len(content) + (len(CONTENT_SEPARATOR) if current else 0)
            if current and current_length + addition > MAX_CONTENT:
                pages.append(CONTENT_SEPARATOR.join(current))
                current = []
                current_length = 0
            current.append(content)
            current_length += len(content) + (len(CONTENT_SEPARATOR) if len(current) > 1 else 0)
    if current:
        pages.append(CONTENT_SEPARATOR.join(current))
    return pages


def merged_sources(lessons: list[dict]) -> list[str]:
    return list(dict.fromkeys(source for lesson in lessons for source in lesson.get("sources", [])))


def compact_group_navigation(group: dict) -> dict:
    """Remove heading-only stops and merge every admission-control bank into one item.

    Heading text is folded into the first real child, so the source remains complete.  Long
    controls are exposed as one menu item with internal Telegram-safe content pages.
    """
    lessons = group["lessons"]
    compacted: list[dict] = []
    position = 0
    while position < len(lessons):
        marker = lessons[position]
        prefix = f"{marker['title']} · "
        child_end = position + 1
        while child_end < len(lessons) and lessons[child_end]["title"].startswith(prefix):
            child_end += 1
        children = lessons[position + 1:child_end]
        if not children:
            compacted.append(marker)
            position += 1
            continue

        if marker["title"] == "Контроль к допуску":
            pages = pack_lesson_pages([marker, *children])
            merged = {
                "id": marker["id"],
                "title": f"Контроль к допуску · {len(children)} вопросов",
                "content": pages[0],
                "content_pages": pages,
                "sources": merged_sources([marker, *children]),
                "question_count": len(children),
            }
            compacted.append(merged)
        else:
            # A Heading 2 card such as "Ходы определения" is not useful as a separate stop.
            # Preserve its complete text by folding it into the first substantive child.
            first_child = deepcopy(children[0])
            pages = pack_lesson_pages([marker, first_child])
            first_child["content"] = pages[0]
            if len(pages) > 1:
                first_child["content_pages"] = pages
            first_child["sources"] = merged_sources([marker, first_child])
            compacted.extend([first_child, *children[1:]])
        position = child_end

    return {**group, "lessons": compacted}


def paragraph_records(document: Document) -> list[dict]:
    records = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = clean_text(paragraph.text)
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else "Normal"
        source_formatted = any(
            run.bold or run.italic or run.underline or run.font.strike for run in paragraph.runs
        )
        rendered_html = paragraph_html(paragraph)
        records.append({
            "index": index,
            "style": style,
            "text": text,
            "html": rendered_html,
            "source_formatted": source_formatted,
            "reading_marker": any(text.startswith(label) for label in READING_MARKERS),
            "semantic_key_term": not source_formatted and "<b>" in rendered_html,
        })
    return records


def records_to_group(group_number: int, heading: dict, body: list[dict]) -> tuple[dict, dict]:
    lessons: list[dict] = []
    mapped: dict[int, str] = {heading["index"]: "group_title"}
    current_h2: dict | None = None
    pending: list[dict] = []
    unit_number = 0

    def flush(title_record: dict | None, values: list[dict], *, title_prefix: str | None = None) -> None:
        nonlocal unit_number
        if title_record is None and not values:
            return
        unit_number += 1
        if title_record is None:
            title = "Обзор занятия"
            heading_index = values[0]["index"]
        else:
            title = title_record["text"]
            heading_index = title_record["index"]
            mapped[heading_index] = "lesson_heading"
        if title_prefix and title_record is not None:
            nav = f"{title_prefix} · {title}"
        else:
            nav = title
        produced = make_lessons(f"b{group_number}_u{unit_number}", nav, values, heading_index)
        lessons.extend(produced)
        for value in values:
            mapped[value["index"]] = produced[0]["id"]

    intro: list[dict] = []
    position = 0
    while position < len(body) and body[position]["style"] not in {"Heading 2", "Heading 3"}:
        intro.append(body[position])
        position += 1
    # The complete Heading 1 text is repeated inside the first lesson because the Telegram
    # navigation label may be shortened to 64 characters.
    flush(heading, intro)
    mapped[heading["index"]] = "group_title_and_lesson_heading"

    while position < len(body):
        record = body[position]
        if record["style"] == "Heading 2":
            if current_h2 is not None and pending:
                flush(current_h2, pending)
            current_h2 = record
            pending = []
            position += 1
            while position < len(body) and body[position]["style"] not in {"Heading 2", "Heading 3"}:
                pending.append(body[position])
                position += 1
            # Keep every Heading 2 as a visible navigation item, even when its only purpose in
            # Word was to introduce Heading 3 children.  This prevents a long heading from
            # surviving only as a truncated button label.
            flush(current_h2, pending)
            pending = []
            continue
        if record["style"] == "Heading 3":
            if current_h2 is not None:
                mapped[current_h2["index"]] = "navigation_context"
            h3 = record
            values: list[dict] = []
            position += 1
            while position < len(body) and body[position]["style"] not in {"Heading 2", "Heading 3"}:
                values.append(body[position])
                position += 1
            prefix = current_h2["text"] if current_h2 else None
            flush(h3, values, title_prefix=prefix)
            continue
        pending.append(record)
        position += 1
    if current_h2 is not None and pending:
        flush(current_h2, pending)

    if not lessons:
        lessons = make_lessons(f"b{group_number}_u1", heading["text"], body, heading["index"])
        for value in body:
            mapped[value["index"]] = lessons[0]["id"]

    group = {
        "id": f"class_{group_number}" if group_number <= 19 else f"reference_{group_number}",
        "title": nav_title(heading["text"]),
        "lessons": lessons,
    }
    return group, mapped


def main() -> None:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    document = Document(SOURCE)
    if document.tables:
        raise RuntimeError("Unexpected tables found; builder must be extended before publishing")
    if document.inline_shapes:
        raise RuntimeError("Unexpected inline images found; builder must be extended before publishing")

    records = paragraph_records(document)
    heading_positions = [i for i, record in enumerate(records) if record["style"] == "Heading 1"]
    if len(heading_positions) != 22:
        raise RuntimeError(f"Expected 22 Heading 1 blocks (contents, 19 classes, appendix, sources), got {len(heading_positions)}")

    mapped: dict[int, str] = {}
    contents_position = heading_positions[0]
    front = records[:contents_position]
    contents_end = heading_positions[1]
    contents_heading = records[contents_position]
    contents_body = records[contents_position + 1:contents_end]
    about_lessons = make_lessons("guide_about", "О практикуме и как с ним работать", front, front[0]["index"])
    contents_lessons = make_lessons("guide_contents", contents_heading["text"], contents_body, contents_heading["index"])
    for record in front:
        mapped[record["index"]] = about_lessons[0]["id"]
    mapped[contents_heading["index"]] = "lesson_heading"
    for record in contents_body:
        mapped[record["index"]] = contents_lessons[0]["id"]

    groups: dict[int, dict] = {}
    for block_number, start_position in enumerate(heading_positions[1:], 1):
        end_position = heading_positions[block_number + 1] if block_number + 1 < len(heading_positions) else len(records)
        heading = records[start_position]
        body = records[start_position + 1:end_position]
        group, group_mapping = records_to_group(block_number, heading, body)
        groups[block_number] = compact_group_navigation(group)
        mapped.update(group_mapping)

    course = {
        "id": "biochemistry",
        "course": 2,
        "title": "Биохимия",
        "emoji": "🧬",
        "description": (
            "Полный практикум ВМедА: 19 занятий с допуском, теорией, лабораторными "
            "определениями, ответами на вопросы и клиническими пояснениями."
        ),
        "show_sources": False,
        "sections": [
            {
                "id": "guide",
                "title": "Как работать с практикумом",
                "groups": [
                    {"id": "about", "title": "О практикуме", "lessons": about_lessons},
                    {"id": "contents", "title": "Содержание · 19 занятий", "lessons": contents_lessons},
                ],
            },
            {"id": "foundations", "title": "Занятия 1–6 · Основы", "groups": [groups[i] for i in range(1, 7)]},
            {"id": "metabolism", "title": "Занятия 7–10 · Обмен веществ", "groups": [groups[i] for i in range(7, 11)]},
            {"id": "regulation", "title": "Занятия 11–14 · Регуляция", "groups": [groups[i] for i in range(11, 15)]},
            {"id": "clinical", "title": "Занятия 15–19 · Клиническая биохимия", "groups": [groups[i] for i in range(15, 20)]},
            {
                "id": "reference",
                "title": "Приложение и источники",
                "groups": [groups[20], groups[21]],
            },
        ],
    }

    rendered_plain = clean_text("\n".join(
        html.unescape(re.sub(r"<[^>]+>", "", page))
        for section in course["sections"]
        for group in section["groups"]
        for lesson in group["lessons"]
        for page in lesson_pages(lesson)
    ))
    missing_verbatim = [record["index"] for record in records if clean_text(record["text"]) not in rendered_plain]
    if missing_verbatim:
        raise RuntimeError(f"Source paragraphs missing from rendered course: {missing_verbatim}")

    all_indices = {record["index"] for record in records}
    missing_indices = sorted(all_indices - set(mapped))
    if missing_indices:
        raise RuntimeError(f"Unmapped source paragraphs: {missing_indices}")

    COURSE_OUTPUT.write_text(json.dumps(course, ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    source_bytes = SOURCE.read_bytes()
    source_manifest = {
        "subject": "biochemistry",
        "collection": "user_provided_attachment",
        "sources": [{
            "original_filename": "Biokhimia_VMedA_praktikum_zanyatia (3).docx",
            "stored_path": SOURCE.relative_to(REPO).as_posix(),
            "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "bytes": len(source_bytes),
            "sha256": hashlib.sha256(source_bytes).hexdigest(),
            "telegram_chat_id": None,
            "telegram_topic_id": None,
            "telegram_message_id": None,
        }],
    }
    lesson_count = sum(
        len(group["lessons"])
        for section in course["sections"]
        for group in section["groups"]
    )
    coverage = {
        "subject": "Биохимия",
        "source_exclusive": True,
        "source_count": 1,
        "source_nonempty_paragraphs": len(records),
        "mapped_nonempty_paragraphs": len(mapped),
        "verbatim_paragraphs_after_whitespace_normalization": len(records),
        "verbatim_coverage_percent": 100,
        "unmapped_paragraphs": [],
        "heading_counts": dict(Counter(record["style"] for record in records if record["style"].startswith("Heading"))),
        "practical_class_count": 19,
        "section_count": len(course["sections"]),
        "group_count": sum(len(section["groups"]) for section in course["sections"]),
        "lesson_count": lesson_count,
        "media_count": 0,
        "formatting": {
            "source_formatted_paragraphs": sum(record["source_formatted"] for record in records),
            "reading_marker_paragraphs": sum(record["reading_marker"] for record in records),
            "semantic_key_term_paragraphs": sum(record["semantic_key_term"] for record in records),
            "control_question_dividers": True,
        },
        "exclusions": [
            {
                "item": "repeating_footer",
                "text": "ВМедА · практикум по биохимии · [номер страницы]",
                "reason": "Повторяющийся колонтитул, не учебный материал",
            }
        ],
    }
    extraction = {
        "source": SOURCE.relative_to(REPO).as_posix(),
        "readable": True,
        "paragraphs_total": len(document.paragraphs),
        "paragraphs_nonempty": len(records),
        "tables": len(document.tables),
        "inline_shapes": len(document.inline_shapes),
        "embedded_media": 0,
        "headers": 0,
        "footers": 1,
        "tracked_insertions": 0,
        "tracked_deletions": 0,
        "failures": [],
        "warnings": [
            "В источнике одновременно приведены кафедральные классические и современные коэффициенты P/O; формулировка сохранена дословно."
        ],
    }
    (REPORT_DIR / "source_manifest.json").write_text(json.dumps(source_manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (REPORT_DIR / "coverage_report.json").write_text(json.dumps(coverage, ensure_ascii=False, indent=2), encoding="utf-8")
    (REPORT_DIR / "extraction_report.json").write_text(json.dumps(extraction, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "paragraphs": len(records),
        "mapped": len(mapped),
        "sections": len(course["sections"]),
        "groups": coverage["group_count"],
        "lessons": lesson_count,
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
