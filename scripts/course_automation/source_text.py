"""Legacy extraction helpers shared by importers; output still requires source review."""
import html
import re
from pathlib import Path

from docx import Document

MAX_CONTENT = 3300


def norm(value: str) -> str:
    value = value.replace("\u00a0", " ").replace("\r", "\n")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def safe(value: str) -> str:
    return html.escape(norm(value), quote=False)


def chunks(text: str, limit: int = MAX_CONTENT) -> list[str]:
    paragraphs = [p.strip() for p in re.split(r"\n+", norm(text)) if p.strip()]
    result, current = [], ""
    for paragraph in paragraphs:
        pieces = [paragraph[i:i + limit] for i in range(0, len(paragraph), limit)]
        for piece in pieces:
            candidate = f"{current}\n{piece}".strip()
            if current and len(candidate) > limit:
                result.append(current)
                current = piece
            else:
                current = candidate
    if current:
        result.append(current)
    return result


def docx_text(path: Path) -> str:
    doc = Document(path)
    values = []
    for block in doc.iter_inner_content():
        if hasattr(block, "text"):
            if norm(block.text):
                values.append(block.text)
        else:
            for row in block.rows:
                value = " | ".join(norm(cell.text) for cell in row.cells)
                if norm(value):
                    values.append(value)
    return norm("\n".join(values))


def docx_paragraphs(path: Path) -> list[str]:
    doc = Document(path)
    return [norm(p.text) for p in doc.paragraphs if norm(p.text)]


def pdf_pages(path: Path) -> list[str]:
    # PDF tooling is an authoring-only dependency. Keep the import local so the
    # production bot and its default CI can validate non-PDF course structures
    # without installing the much larger automation dependency set.
    from pypdf import PdfReader

    return [norm(page.extract_text() or "") for page in PdfReader(path).pages]


def lesson(lesson_id: str, title: str, body: str, source: str, media=None) -> dict:
    result = {
        "id": lesson_id,
        "title": norm(title)[:180],
        "content": f"<b>Материал курса</b>\n\n{safe(body)}",
        "sources": [source],
    }
    if media:
        result["media"] = media
    return result


def page_lessons(prefix: str, pages: list[str], source_name: str, label: str) -> list[dict]:
    result = []
    for page_no, text in enumerate(pages, 1):
        if len(text) < 40:
            continue
        page_chunks = chunks(text)
        first = re.sub(r"^\[.*?\]\s*", "", text).split("\n", 1)[0][:100]
        for part, body in enumerate(page_chunks, 1):
            suffix = f", часть {part}" if len(page_chunks) > 1 else ""
            result.append(lesson(
                f"{prefix}_p{page_no}_{part}",
                f"{label}: стр. {page_no}{suffix} — {first}", body,
                f"{source_name}, стр. {page_no}",
            ))
    return result


def numbered_lessons(prefix: str, text: str, source_name: str, label: str) -> list[dict]:
    pattern = re.compile(r"(?m)^(?:(?:ВОПРОС|БИЛЕТ)\s*№?\s*)?(\d{1,3})[\).!:–-]?\s*(?=\S)", re.I)
    matches = list(pattern.finditer(text))
    if len(matches) < 3:
        return [lesson(f"{prefix}_{i}", f"{label}, часть {i}", body, source_name)
                for i, body in enumerate(chunks(text), 1)]
    result = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = norm(text[match.start():end])
        number = match.group(1)
        title_line = block.split("\n", 1)[0][:130]
        for part, body in enumerate(chunks(block), 1):
            extra = f", часть {part}" if len(chunks(block)) > 1 else ""
            result.append(lesson(f"{prefix}_{index + 1}_{part}", f"{label} {number}{extra}: {title_line}", body, source_name))
    return result


def ranged_question_lessons(prefix: str, text: str, source_name: str, label: str, numbers: range, marker: str = r"[)!]") -> list[dict]:
    wanted = set(numbers)
    matches = [match for match in re.finditer(rf"(?m)^(\d{{1,3}}){marker}\s*(?=\S)", text)
               if int(match.group(1)) in wanted]
    # A number can recur inside an answer; the first ordered occurrence is the question heading.
    selected, last = [], -1
    for number in numbers:
        match = next((item for item in matches if int(item.group(1)) == number and item.start() > last), None)
        if match:
            selected.append(match)
            last = match.start()
    result = []
    for index, match in enumerate(selected):
        end = selected[index + 1].start() if index + 1 < len(selected) else len(text)
        block = norm(text[match.start():end])
        parts = chunks(block)
        for part, body in enumerate(parts, 1):
            suffix = f", часть {part}" if len(parts) > 1 else ""
            result.append(lesson(f"{prefix}_{match.group(1)}_{part}", f"{label} {match.group(1)}{suffix}: {block.splitlines()[0][:130]}", body, source_name))
    return result


def multiple_choice_lessons(prefix: str, path: Path, label: str) -> list[dict]:
    values = docx_paragraphs(path)
    option = re.compile(r"^[а-яёa-z][).]\s*", re.I)
    groups, current = [], []
    for value in values:
        if option.match(value):
            if current:
                current.append(value)
            continue
        if current:
            groups.append(current)
        current = [value]
    if current:
        groups.append(current)
    result = []
    for number, group in enumerate(groups, 1):
        body = "\n".join(group)
        if len(body) < 20:
            continue
        result.append(lesson(f"{prefix}_{number}", f"{label} {number}: {group[0][:135]}", body, path.name))
    return result
